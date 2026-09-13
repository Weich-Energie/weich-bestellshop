"""Baut die Arbeitsmappe zu den geratenen R+F-Zuordnungen."""
import io
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

QUELLE, ZIEL = sys.argv[1], sys.argv[2]
daten = json.load(io.open(QUELLE, encoding="utf-8"))


def klar(s):
    s = str(s or "")
    for a, b in (("&quot;", '"'), ("&amp;", "&"), ("&lt;", "<"), ("&gt;", ">")):
        s = s.replace(a, b)
    return " ".join(s.split())


def zahl(x):
    if x is None or x == "":
        return None
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


SCHRIFT = "Arial"
KOPF_FUELL = PatternFill("solid", fgColor="1F3B4D")
EINGABE_FUELL = PatternFill("solid", fgColor="FFF2CC")
WARN_FUELL = PatternFill("solid", fgColor="FCE4E4")
TITEL_F = Font(name=SCHRIFT, size=14, bold=True, color="1F3B4D")
KOPF_F = Font(name=SCHRIFT, size=10, bold=True, color="FFFFFF")
TEXT_F = Font(name=SCHRIFT, size=10)
FETT_F = Font(name=SCHRIFT, size=10, bold=True)
LEISE_F = Font(name=SCHRIFT, size=9, color="5D6871")
duenn = Side(style="thin", color="D3D9DD")
RAHMEN = Border(left=duenn, right=duenn, top=duenn, bottom=duenn)

wb = Workbook()

# ── Blatt 1: Übersicht und Anleitung ──────────────────────────────────────
ue = wb.active
ue.title = "Übersicht"
ue.sheet_view.showGridLines = False

ue["A1"] = "Geratene R+F-Zuordnungen"
ue["A1"].font = TITEL_F
ue["A2"] = ("91 GUT-Artikel, deren R+F-Entsprechung per Textsuche bestimmt wurde "
            "und nicht durch Rechnung, Lagerauszug oder Entscheidung belegt ist.")
ue["A2"].font = TEXT_F
ue["A3"] = "Stand 13.09.2026 · Weich Energie · Aufmaß-Projekt"
ue["A3"].font = LEISE_F

ue["A5"] = "So wird die Liste ausgefüllt"
ue["A5"].font = FETT_F
anleitung = [
    "Im Blatt „Zuordnungen“ sind nur die beiden gelben Spalten zu füllen:",
    "    Spalte L  „richtige R+F-Nummer“   die Artikelnummer, die stattdessen gilt",
    "    Spalte M  „Bemerkung“             etwa „wird nicht mehr verbaut“ oder „Preis prüfen“",
    "Alle übrigen Spalten bitte stehen lassen, sie sind die Grundlage für den Abgleich.",
    "",
    "Spalte K sagt, ob der zugeordnete R+F-Artikel überhaupt im Lagerauszug steht.",
    "Steht dort „nein“, lässt sich der Preis an nichts prüfen — dort lagen alle",
    "schweren Fehler, die bisher gefunden wurden.",
]
for i, t in enumerate(anleitung):
    ue.cell(row=6 + i, column=1, value=t).font = TEXT_F

ue["A15"] = "Beispielzeile"
ue["A15"].font = FETT_F
ue["A16"] = "GUT-Nummer"
ue["B16"] = "richtige R+F-Nummer"
ue["C16"] = "Bemerkung"
for c in "ABC":
    ue[c + "16"].font = KOPF_F
    ue[c + "16"].fill = KOPF_FUELL
    ue[c + "16"].border = RAHMEN
ue["A17"] = "POVUS3525A"
ue["B17"] = "7020900559045"
ue["C17"] = "Prestabo läuft aus, Preis vom 28er übernommen"
for c in "ABC":
    ue[c + "17"].font = TEXT_F
    ue[c + "17"].border = RAHMEN
ue["B17"].fill = EINGABE_FUELL
ue["C17"].fill = EINGABE_FUELL
ue["A18"] = "Nur ein Beispiel für das Format — diese Zeile bitte nicht übertragen."
ue["A18"].font = LEISE_F

# Kennzahlen je Kategorie, als Formeln auf das Datenblatt
ue["A21"] = "Nach Kategorie"
ue["A21"].font = FETT_F
kopfzeile = ["Kategorie", "Artikel", "Warenwert", "davon ohne Regal-Gegenstück"]
for j, t in enumerate(kopfzeile, start=1):
    z = ue.cell(row=22, column=j, value=t)
    z.font = KOPF_F
    z.fill = KOPF_FUELL
    z.border = RAHMEN

kategorien = []
for d in daten:
    k = d.get("kategorie") or "ohne Angabe"
    if k not in kategorien:
        kategorien.append(k)
kategorien.sort()

n = len(daten)
ende = 1 + n  # letzte Datenzeile im Blatt "Zuordnungen"
for i, k in enumerate(kategorien):
    r = 23 + i
    ue.cell(row=r, column=1, value=k).font = TEXT_F
    ue.cell(row=r, column=2,
            value=f'=COUNTIF(Zuordnungen!$A$2:$A${ende},$A{r})')
    ue.cell(row=r, column=3,
            value=f'=SUMIF(Zuordnungen!$A$2:$A${ende},$A{r},Zuordnungen!$H$2:$H${ende})')
    ue.cell(row=r, column=4,
            value=f'=COUNTIFS(Zuordnungen!$A$2:$A${ende},$A{r},'
                  f'Zuordnungen!$K$2:$K${ende},"nein")')
    for j in (2, 3, 4):
        ue.cell(row=r, column=j).font = TEXT_F
    ue.cell(row=r, column=3).number_format = '#,##0 "€"'
    for j in range(1, 5):
        ue.cell(row=r, column=j).border = RAHMEN

summe = 23 + len(kategorien)
ue.cell(row=summe, column=1, value="Summe").font = FETT_F
for j, sp in ((2, "B"), (3, "C"), (4, "D")):
    z = ue.cell(row=summe, column=j,
                value=f"=SUM({sp}23:{sp}{summe - 1})")
    z.font = FETT_F
    z.border = RAHMEN
ue.cell(row=summe, column=1).border = RAHMEN
ue.cell(row=summe, column=3).number_format = '#,##0 "€"'

for sp, breite in (("A", 46), ("B", 22), ("C", 46), ("D", 30)):
    ue.column_dimensions[sp].width = breite

# ── Blatt 2: die Zuordnungen ──────────────────────────────────────────────
ws = wb.create_sheet("Zuordnungen")
spalten = [
    ("Kategorie", 16, None),
    ("Formteil-Gruppe", 34, None),
    ("GUT-Nummer", 15, None),
    ("Bezeichnung", 52, None),
    ("Zusatz", 26, None),
    ("Menge", 9, "#,##0"),
    ("zugeordneter EK", 15, '#,##0.00 "€"'),
    ("Warenwert", 12, '#,##0 "€"'),
    ("zugeordnete R+F-Nr", 19, None),
    ("R+F-Bezeichnung", 44, None),
    ("R+F im Lagerauszug", 18, None),
    ("richtige R+F-Nummer", 21, None),
    ("Bemerkung", 40, None),
]
for j, (t, breite, _) in enumerate(spalten, start=1):
    z = ws.cell(row=1, column=j, value=t)
    z.font = KOPF_F
    z.fill = KOPF_FUELL
    z.border = RAHMEN
    z.alignment = Alignment(vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(j)].width = breite
ws.row_dimensions[1].height = 30

for i, d in enumerate(daten, start=2):
    werte = [
        d.get("kategorie") or "ohne Angabe",
        d.get("gruppe") or "",
        d.get("gut_nr"),
        klar(d.get("text")),
        klar(d.get("text2")),
        zahl(d.get("menge")),
        zahl(d.get("rf_ek")),
        zahl(d.get("wert")),
        d.get("rf_nr") or "",
        klar(d.get("rf_name")),
        "nein" if d.get("ohne_regal") else "ja",
        None,
        None,
    ]
    for j, w in enumerate(werte, start=1):
        z = ws.cell(row=i, column=j, value=w)
        z.font = TEXT_F
        z.border = RAHMEN
        z.alignment = Alignment(vertical="top", wrap_text=(j in (4, 5, 10, 13)))
        fmt = spalten[j - 1][2]
        if fmt:
            z.number_format = fmt
    if werte[10] == "nein":
        ws.cell(row=i, column=11).fill = WARN_FUELL
    ws.cell(row=i, column=12).fill = EINGABE_FUELL
    ws.cell(row=i, column=13).fill = EINGABE_FUELL

ws.freeze_panes = "D2"
ws.auto_filter.ref = f"A1:M{1 + n}"

pruef = DataValidation(type="textLength", operator="lessThanOrEqual",
                       formula1="14", allow_blank=True,
                       error="Eine R+F-Artikelnummer hat höchstens 14 Stellen.",
                       errorTitle="Artikelnummer prüfen")
ws.add_data_validation(pruef)
pruef.add(f"L2:L{1 + n}")

wb.save(ZIEL)
print(f"{n} Zeilen, {len(kategorien)} Kategorien -> {ZIEL}")
